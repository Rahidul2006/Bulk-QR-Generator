import csv
import io
import os
import re
import threading
import zipfile
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
import tkinter as tk
from typing import Any, cast

import openpyxl
import qrcode
from qrcode import constants as qr_constants
from tkinterdnd2 import DND_FILES, TkinterDnD  # type: ignore[reportMissingImports]
from PIL import Image, ImageTk


PLACEHOLDER_RE = re.compile(r"{{\s*([^{}]+?)\s*}}")
INVALID_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')
ERROR_CORRECTION = {
	"Low": qr_constants.ERROR_CORRECT_L,
	"Medium": qr_constants.ERROR_CORRECT_M,
	"Quartile": qr_constants.ERROR_CORRECT_Q,
	"High": qr_constants.ERROR_CORRECT_H,
}


@dataclass
class GeneratedQR:
	row_number: int
	filename: str
	url: str
	image_bytes: bytes
	status: str = "Success"
	reason: str = ""


def clean_filename(value, fallback):
	cleaned = INVALID_FILENAME_RE.sub("_", str(value or "")).strip().strip(".")
	cleaned = re.sub(r"\s+", "_", cleaned)
	return cleaned or fallback


def read_workbook(filename):
	workbook = openpyxl.load_workbook(filename, read_only=True, data_only=True)
	sheet = workbook.active
	if sheet is None:
		workbook.close()
		raise ValueError("The workbook does not contain an active worksheet.")
	rows = list(sheet.iter_rows(values_only=True))
	workbook.close()
	if not rows:
		raise ValueError("The workbook does not contain a header row.")
	headers = [str(value).strip() if value is not None else "" for value in rows[0]]
	headers = [header for header in headers if header]
	if not headers:
		raise ValueError("The workbook header row is empty.")
	data = []
	for values in rows[1:]:
		record = {header: (values[index] if index < len(values) else "") for index, header in enumerate(headers)}
		if any(value not in (None, "") for value in record.values()):
			data.append(record)
	if not data:
		raise ValueError("The workbook contains headers but no data rows.")
	return headers, data


def placeholders(template):
	return list(dict.fromkeys(PLACEHOLDER_RE.findall(template)))


def make_url(template, row):
	def replace(match):
		column = match.group(1).strip()
		value = row.get(column, "")
		from urllib.parse import quote
		return quote(str(value if value is not None else ""), safe="")
	return PLACEHOLDER_RE.sub(replace, template)


def qr_bytes(url, size, correction, output_format):
	qr = qrcode.QRCode(error_correction=ERROR_CORRECTION[correction], box_size=10, border=4)
	qr.add_data(url)
	qr.make(fit=True)
	if output_format == "SVG":
		from qrcode.image.svg import SvgPathImage
		stream = io.BytesIO()
		qr.make_image(image_factory=SvgPathImage).save(stream)
		return stream.getvalue()
	image = cast(Image.Image, qr.make_image(fill_color="black", back_color="white")).convert("RGB")
	image = image.resize((size, size), Image.Resampling.LANCZOS)
	stream = io.BytesIO()
	image.save(stream, format="PNG")
	return stream.getvalue()


def generate_rows(rows, headers, template, filename_column, prefix, suffix, size, correction, output_format):
	unknown = [column for column in placeholders(template) if column not in headers]
	if unknown:
		raise ValueError("Unknown column(s): " + ", ".join("{{" + column + "}}" for column in unknown))
	used = set()
	results = []
	for index, row in enumerate(rows, 1):
		try:
			url = make_url(template, row)
			if not url.strip() or not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", url):
				raise ValueError("The generated URL is not valid.")
			base = clean_filename(row.get(filename_column), f"QR_{index:03d}")
			extension = output_format.lower()
			candidate = f"{prefix}{base}{suffix}.{extension}"
			stem, ext = os.path.splitext(candidate)
			count = 2
			while candidate.lower() in used:
				candidate = f"{stem}_{count}{ext}"
				count += 1
			used.add(candidate.lower())
			results.append(GeneratedQR(index, candidate, url, qr_bytes(url, size, correction, output_format)))
		except Exception as error:
			results.append(GeneratedQR(index, "", "", b"", "Failed", str(error)))
	return results


class QRGeneratorApp:
	def __init__(self, root):
		self.root = root
		self.root.title("Bulk QR Studio")
		self.root.geometry("1180x820")
		self.root.minsize(900, 650)
		self.root.configure(bg="#f4f7fb")
		self.headers = []
		self.rows = []
		self.results = []
		self.preview_photo = None
		self.build_style()
		self.build_ui()

	def build_style(self):
		style = ttk.Style()
		style.theme_use("clam")
		style.configure("TFrame", background="#f4f7fb")
		style.configure("Card.TFrame", background="white")
		style.configure("TLabel", background="#f4f7fb", foreground="#24324a", font=("Segoe UI", 10))
		style.configure("Card.TLabel", background="white")
		style.configure("Title.TLabel", background="#f4f7fb", foreground="#12213d", font=("Segoe UI", 25, "bold"))
		style.configure("Subtitle.TLabel", background="#f4f7fb", foreground="#66758f", font=("Segoe UI", 11))
		style.configure("Section.TLabel", background="white", foreground="#12213d", font=("Segoe UI", 14, "bold"))
		style.configure("TButton", font=("Segoe UI", 10), padding=(12, 8))
		style.configure("Primary.TButton", background="#1264d6", foreground="white", font=("Segoe UI", 11, "bold"))
		style.map("Primary.TButton", background=[("active", "#0b4fae")])
		style.configure("Treeview", rowheight=28, font=("Segoe UI", 9))
		style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))

	def build_ui(self):
		outer = ttk.Frame(self.root, padding=26)
		outer.pack(fill="both", expand=True)
		ttk.Label(outer, text="Bulk QR Studio", style="Title.TLabel").pack(anchor="w")
		ttk.Label(outer, text="Turn every row in an Excel sheet into a personalized QR code.", style="Subtitle.TLabel").pack(anchor="w", pady=(2, 18))
		self.status = ttk.Label(outer, text="01 Upload  →  02 URL  →  03 Naming  →  04 Generate  →  05 Download", style="Subtitle.TLabel")
		self.status.pack(anchor="w", pady=(0, 14))

		content = ttk.Frame(outer)
		content.pack(fill="both", expand=True)
		left_shell = ttk.Frame(content, style="Card.TFrame")
		left_shell.pack(side="left", fill="y", padx=(0, 14))
		left_canvas = tk.Canvas(left_shell, width=350, bg="white", highlightthickness=0)
		left_scrollbar = ttk.Scrollbar(left_shell, orient="vertical", command=left_canvas.yview)
		left_canvas.configure(yscrollcommand=left_scrollbar.set)
		left_canvas.pack(side="left", fill="y", expand=True)
		left_scrollbar.pack(side="right", fill="y")
		left = ttk.Frame(left_canvas, style="Card.TFrame", padding=20)
		left_window = left_canvas.create_window((0, 0), window=left, anchor="nw", width=350)
		left.bind("<Configure>", lambda _: left_canvas.configure(scrollregion=left_canvas.bbox("all")))
		left_canvas.bind("<Configure>", lambda event: left_canvas.itemconfigure(left_window, width=event.width))
		left_canvas.bind_all("<MouseWheel>", lambda event: left_canvas.yview_scroll(int(-event.delta / 120), "units"))
		right = ttk.Frame(content, style="Card.TFrame", padding=20)
		right.pack(side="left", fill="both", expand=True)
		self.build_controls(left)
		self.build_results(right)

	def build_controls(self, parent):
		ttk.Label(parent, text="1. Upload Excel", style="Section.TLabel").pack(anchor="w")
		self.drop = tk.Label(parent, text="Drop a .xlsx file here\nor browse to choose one", bg="#eef5ff", fg="#1264d6", width=34, height=4, font=("Segoe UI", 11, "bold"), relief="solid", bd=1, cursor="hand2")
		self.drop.pack(fill="x", pady=(12, 8))
		self.drop.bind("<Button-1>", lambda _: self.choose_file())
		drop_target = cast(Any, self.drop)
		drop_target.drop_target_register(DND_FILES)
		drop_target.dnd_bind("<<Drop>>", self.drop_file)
		ttk.Button(parent, text="Browse .xlsx file", command=self.choose_file).pack(fill="x")
		self.file_label = ttk.Label(parent, text="No file selected", style="Card.TLabel", wraplength=300)
		self.file_label.pack(anchor="w", pady=(8, 12))
		self.column_label = ttk.Label(parent, text="Detected columns will appear here.", style="Card.TLabel", wraplength=300)
		self.column_label.pack(anchor="w")

		ttk.Separator(parent).pack(fill="x", pady=18)
		ttk.Label(parent, text="2. Configure URL", style="Section.TLabel").pack(anchor="w")
		ttk.Label(parent, text="URL template", style="Card.TLabel").pack(anchor="w", pady=(12, 4))
		self.template = tk.Text(parent, height=3, width=38, font=("Segoe UI", 10), relief="solid", bd=1, wrap="word")
		self.template.pack(fill="x")
		self.template.bind("<KeyRelease>", lambda _: self.update_preview())
		ttk.Label(parent, text="Insert column", style="Card.TLabel").pack(anchor="w", pady=(10, 4))
		self.columns_frame = ttk.Frame(parent, style="Card.TFrame")
		self.columns_frame.pack(fill="x")
		self.example = ttk.Label(parent, text="Example generated URL: upload a workbook first.", style="Card.TLabel", wraplength=300)
		self.example.pack(anchor="w", pady=(12, 0))
		self.validation = ttk.Label(parent, text="", style="Card.TLabel", foreground="#c0392b", wraplength=300)
		self.validation.pack(anchor="w", pady=(5, 0))

		ttk.Separator(parent).pack(fill="x", pady=18)
		ttk.Label(parent, text="3. QR naming & settings", style="Section.TLabel").pack(anchor="w")
		self.filename_column = tk.StringVar()
		self.filename_menu = ttk.Combobox(parent, textvariable=self.filename_column, state="readonly")
		self.filename_menu.pack(fill="x", pady=(12, 6))
		self.prefix = self.field(parent, "Filename prefix")
		self.suffix = self.field(parent, "Filename suffix")
		settings = ttk.Frame(parent, style="Card.TFrame")
		settings.pack(fill="x", pady=(4, 0))
		self.size = tk.StringVar(value="256")
		self.format = tk.StringVar(value="PNG")
		self.correction = tk.StringVar(value="Medium")
		for label, variable, values in (("Size", self.size, ("256", "512", "1024")), ("Format", self.format, ("PNG", "SVG")), ("Error correction", self.correction, tuple(ERROR_CORRECTION))):
			ttk.Label(settings, text=label, style="Card.TLabel").pack(anchor="w", pady=(5, 2))
			ttk.Combobox(settings, textvariable=variable, values=values, state="readonly").pack(fill="x")
		self.generate_button = ttk.Button(parent, text="Generate QR Codes", style="Primary.TButton", command=self.start_generation)
		self.generate_button.pack(fill="x", pady=(20, 0))
		self.progress = ttk.Progressbar(parent, mode="determinate")
		self.progress.pack(fill="x", pady=(10, 3))
		self.progress_label = ttk.Label(parent, text="", style="Card.TLabel")
		self.progress_label.pack(anchor="w")

	def field(self, parent, label):
		ttk.Label(parent, text=label, style="Card.TLabel").pack(anchor="w", pady=(7, 2))
		variable = tk.StringVar()
		ttk.Entry(parent, textvariable=variable).pack(fill="x")
		return variable

	def build_results(self, parent):
		ttk.Label(parent, text="Excel preview", style="Section.TLabel").pack(anchor="w")
		self.preview_tree = ttk.Treeview(parent, show="headings", height=7)
		self.preview_tree.pack(fill="x", pady=(12, 4))
		self.preview_count = ttk.Label(parent, text="Upload a workbook to preview the first 10 rows.", style="Card.TLabel")
		self.preview_count.pack(anchor="w")
		ttk.Separator(parent).pack(fill="x", pady=18)
		ttk.Label(parent, text="Generation results", style="Section.TLabel").pack(anchor="w")
		self.summary = ttk.Label(parent, text="No QR codes generated yet.", style="Card.TLabel")
		self.summary.pack(anchor="w", pady=(8, 8))
		actions = ttk.Frame(parent, style="Card.TFrame")
		actions.pack(fill="x")
		self.zip_button = ttk.Button(actions, text="Download All as ZIP", command=self.download_zip, state="disabled")
		self.zip_button.pack(side="left")
		self.single_button = ttk.Button(actions, text="Download Selected", command=self.download_selected, state="disabled")
		self.single_button.pack(side="left", padx=8)
		self.report_button = ttk.Button(actions, text="Download Report CSV", command=self.download_report, state="disabled")
		self.report_button.pack(side="left")
		body = ttk.Frame(parent, style="Card.TFrame")
		body.pack(fill="both", expand=True, pady=(12, 0))
		self.image_preview = ttk.Label(body, text="Select a generated row to preview its QR image.", style="Card.TLabel", anchor="center", width=34)
		self.image_preview.pack(side="right", fill="y", padx=(16, 0))
		self.result_tree = ttk.Treeview(body, columns=("row", "filename", "url", "status"), show="headings")
		for column, heading, width in (("row", "Row", 50), ("filename", "Filename", 180), ("url", "Generated URL", 390), ("status", "Status", 110)):
			self.result_tree.heading(column, text=heading)
			self.result_tree.column(column, width=width, anchor="w")
		self.result_tree.pack(side="left", fill="both", expand=True)
		self.result_tree.bind("<<TreeviewSelect>>", self.show_selected)
		scroll = ttk.Scrollbar(body, orient="vertical", command=self.result_tree.yview)
		scroll.pack(side="right", fill="y")
		self.result_tree.configure(yscrollcommand=scroll.set)

	def choose_file(self):
		filename = filedialog.askopenfilename(filetypes=[("Excel workbook", "*.xlsx")])
		if not filename:
			return
		self.load_file(filename)

	def drop_file(self, event):
		filename = self.root.tk.splitlist(event.data)[0]
		if filename.lower().endswith(".xlsx"):
			self.load_file(filename)
		else:
			messagebox.showerror("Invalid file", "Only .xlsx Excel files are supported.")

	def load_file(self, filename):
		try:
			self.headers, self.rows = read_workbook(filename)
		except Exception as error:
			self.headers, self.rows = [], []
			messagebox.showerror("Invalid Excel file", str(error))
			return
		self.file_label.configure(text=f"{Path(filename).name}  |  {len(self.rows)} rows")
		self.column_label.configure(text="Detected columns: " + ", ".join(self.headers))
		self.filename_menu["values"] = self.headers
		self.filename_column.set(self.headers[0])
		for child in self.columns_frame.winfo_children():
			child.destroy()
		for header in self.headers:
			ttk.Button(self.columns_frame, text=header, command=lambda value=header: self.insert_column(value)).pack(side="left", padx=(0, 4), pady=2)
		self.populate_preview()
		self.update_preview()

	def populate_preview(self):
		self.preview_tree.delete(*self.preview_tree.get_children())
		self.preview_tree["columns"] = self.headers
		for header in self.headers:
			self.preview_tree.heading(header, text=header)
			self.preview_tree.column(header, width=max(90, min(180, len(header) * 10)))
		for row in self.rows[:10]:
			self.preview_tree.insert("", "end", values=[str(row.get(header, "")) for header in self.headers])
		self.preview_count.configure(text=f"Showing {min(10, len(self.rows))} of {len(self.rows)} rows")

	def insert_column(self, column):
		self.template.insert(tk.INSERT, "{{" + column + "}}")
		self.update_preview()

	def update_preview(self):
		template = self.template.get("1.0", "end-1c")
		if not self.rows or not template.strip():
			self.example.configure(text="Example generated URL: upload a workbook and enter a template.")
			self.validation.configure(text="")
			return
		unknown = [column for column in placeholders(template) if column not in self.headers]
		if unknown:
			message = "Unknown column: " + ", ".join("{{" + column + "}}" for column in unknown)
			self.validation.configure(text=message)
		else:
			self.validation.configure(text="")
		self.example.configure(text="Example generated URL:\n" + make_url(template, self.rows[0]))

	def start_generation(self):
		template = self.template.get("1.0", "end-1c").strip()
		if not self.rows:
			messagebox.showerror("Upload required", "Choose a valid .xlsx file with at least one data row.")
			return
		if not template:
			messagebox.showerror("URL required", "Enter a URL template before generating.")
			return
		if not self.filename_column.get():
			messagebox.showerror("Filename column required", "Select the Excel column used for filenames.")
			return
		unknown = [column for column in placeholders(template) if column not in self.headers]
		if unknown:
			messagebox.showerror("Unknown column", "The placeholder(s) do not match Excel columns: " + ", ".join(unknown))
			return
		self.generate_button.configure(state="disabled")
		self.progress.configure(value=0, maximum=len(self.rows))
		self.progress_label.configure(text=f"Generating 0 / {len(self.rows)}")
		settings = (template, self.filename_column.get(), self.prefix.get(), self.suffix.get(), int(self.size.get()), self.correction.get(), self.format.get())
		threading.Thread(target=self.generate_worker, args=(settings,), daemon=True).start()

	def generate_worker(self, settings):
		try:
			self.results = generate_rows(self.rows, self.headers, *settings)
			self.root.after(0, self.finish_generation)
		except Exception as error:
			self.root.after(0, lambda: messagebox.showerror("Generation failed", str(error)))
			self.root.after(0, lambda: self.generate_button.configure(state="normal"))

	def finish_generation(self):
		self.result_tree.delete(*self.result_tree.get_children())
		successful = sum(result.status == "Success" for result in self.results)
		for result in self.results:
			status = result.status if result.status == "Success" else f"Failed: {result.reason}"
			self.result_tree.insert("", "end", values=(result.row_number, result.filename or "-", result.url or "-", status))
		self.summary.configure(text=f"{len(self.results)} total rows  |  {successful} generated  |  {len(self.results) - successful} failed")
		self.progress.configure(value=len(self.results))
		self.progress_label.configure(text=f"Generation complete: {successful} / {len(self.results)} generated")
		self.generate_button.configure(state="normal")
		state = "normal" if successful else "disabled"
		self.zip_button.configure(state=state)
		self.single_button.configure(state=state)
		self.report_button.configure(state="normal")

	def show_selected(self, _event=None):
		selected = self.result_tree.selection()
		if not selected:
			return
		result = self.results[self.result_tree.index(selected[0])]
		if result.status != "Success":
			self.image_preview.configure(image="", text=result.reason)
			self.preview_photo = None
			return
		if self.format.get() == "SVG":
			self.image_preview.configure(image="", text="SVG generated\nUse Download Selected to save it.")
			self.preview_photo = None
			return
		image = Image.open(io.BytesIO(result.image_bytes))
		image.thumbnail((220, 220))
		self.preview_photo = ImageTk.PhotoImage(image)
		self.image_preview.configure(image=self.preview_photo, text=result.filename)

	def download_selected(self):
		selected = self.result_tree.selection()
		if not selected:
			return
		result = self.results[self.result_tree.index(selected[0])]
		if result.status != "Success":
			return
		filename = filedialog.asksaveasfilename(initialfile=result.filename, defaultextension="." + self.format.get().lower(), filetypes=[(self.format.get() + " image", "*." + self.format.get().lower())])
		if filename:
			with open(filename, "wb") as output:
				output.write(result.image_bytes)

	def download_zip(self):
		filename = filedialog.asksaveasfilename(initialfile="QR_Codes.zip", defaultextension=".zip", filetypes=[("ZIP archive", "*.zip")])
		if not filename:
			return
		with zipfile.ZipFile(filename, "w", zipfile.ZIP_DEFLATED) as archive:
			for result in self.results:
				if result.status == "Success":
					archive.writestr(result.filename, result.image_bytes)
		messagebox.showinfo("ZIP ready", f"Saved {sum(result.status == 'Success' for result in self.results)} QR codes.")

	def download_report(self):
		filename = filedialog.asksaveasfilename(initialfile="QR_generation_report.csv", defaultextension=".csv", filetypes=[("CSV report", "*.csv")])
		if not filename:
			return
		with open(filename, "w", newline="", encoding="utf-8-sig") as report:
			writer = csv.writer(report)
			writer.writerow(("Row", "Filename", "Generated URL", "Status", "Reason"))
			for result in self.results:
				writer.writerow((result.row_number, result.filename, result.url, result.status, result.reason))
		messagebox.showinfo("Report ready", "The generation report was saved.")


if __name__ == "__main__":
	application = TkinterDnD.Tk()
	QRGeneratorApp(application)
	application.mainloop()